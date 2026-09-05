"""
Preisliste Schwabenprint (Excel) -> JSON-Zeilen auf stdout.

Zwei Sheet-Familien:
  matrix    Wochen-/Wand-/Tisch-/Wochentischkalender
            Kopf: 3 Dim-Zeilen (Format / Blatt[/Farbigkeit] / Sorte) ueber der
            Zeile mit "Auflage" bzw. "Kombi" in Spalte A, die zugleich die
            kombinierten Spalten-Keys traegt. Danach je Zeile: Auflage | Preise.
  longform  Speisekarte
            Kopfzeile mit NUMMER | Format | GROESSE | FARBE | ... | 10 | 20 | ...
            je Datenzeile ein Artikel, Preise ueber die Auflage-Spalten.

Nicht behandelt (eigene Logik noetig): Spiralbooklet, Bloecke (Staffel/"weitere N").
"""
import json
import sys

import openpyxl

MATRIX = {
    "Preisliste Wochenkalender": {"kategorie": "Wochenkalender", "gruppe": "WOK"},
    "Preisliste Wandkalender": {"kategorie": "Wandkalender", "gruppe": "DWK"},
    "Preisliste Tischkalender": {"kategorie": "Tischkalender", "gruppe": "DKL"},
    "Preisliste Wochentischkalender": {"kategorie": "Wochentischkalender", "gruppe": "DWT"},
}
LONGFORM = {
    "Preisliste Speisekarte": {"kategorie": "Speisekarte", "gruppe": "DSK"},
}


def s(v):
    return "" if v is None else str(v).strip()


def as_int(v):
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None


def num(v):
    try:
        return round(float(v), 4)
    except (TypeError, ValueError):
        return None


def parse_matrix(ws, cfg):
    rows = list(ws.iter_rows(values_only=True))
    key_row = None
    for i, r in enumerate(rows):
        if r and s(r[0]).lower() in ("auflage", "kombi"):
            key_row = i
            break
    if key_row is None or key_row < 3:
        return []
    fmt_r, blatt_r, sorte_r = rows[key_row - 3], rows[key_row - 2], rows[key_row - 1]
    keys = rows[key_row]

    cols = []
    for c in range(1, len(keys)):
        key = s(keys[c])
        if not key:
            continue
        blatt_cell = s(blatt_r[c]) if c < len(blatt_r) else ""
        farb = ""
        blatt = as_int(blatt_cell.split()[0]) if blatt_cell else None
        if "/" in blatt_cell:
            farb = blatt_cell.split()[-1]
        cols.append(
            {
                "col": c,
                "spalten_key": key,
                "format": s(fmt_r[c]) if c < len(fmt_r) else "",
                "blatt": blatt,
                "farbigkeit": farb,
                "sorte": s(sorte_r[c]) if c < len(sorte_r) else "",
            }
        )

    out = []
    for r in rows[key_row + 1 :]:
        aufl = as_int(r[0]) if r else None
        if aufl is None:
            continue
        for co in cols:
            p = num(r[co["col"]]) if co["col"] < len(r) else None
            if p is None:
                continue
            out.append(
                {
                    "kategorie": cfg["kategorie"],
                    "produktgruppe": cfg["gruppe"],
                    "format": co["format"] or None,
                    "blatt": co["blatt"],
                    "sorte": co["sorte"] or None,
                    "farbigkeit": co["farbigkeit"] or None,
                    "spalten_key": co["spalten_key"],
                    "auflage": aufl,
                    "preis_netto": p,
                }
            )
    return out


def parse_longform(ws, cfg):
    rows = list(ws.iter_rows(values_only=True))
    hdr = None
    for i, r in enumerate(rows):
        if r and any(s(x).upper() == "NUMMER" for x in r):
            hdr = i
            break
    if hdr is None:
        return []
    head = rows[hdr]
    idx = {s(x).upper(): j for j, x in enumerate(head) if s(x)}
    col_nr = idx.get("NUMMER")
    col_fmt = idx.get("FORMAT")
    col_gr = idx.get("GROESSE") or idx.get("GRÖSSE")
    col_farbe = idx.get("FARBE")
    # Auflage-Spalten: numerische Header rechts der Stammspalten
    aufl_cols = [(j, as_int(x)) for j, x in enumerate(head) if as_int(x) is not None and j > (col_farbe or 0)]
    cello_row = rows[hdr - 1] if hdr > 0 else ()

    out = []
    for r in rows[hdr + 1 :]:
        if not r or col_nr is None or not s(r[col_nr]):
            continue
        nr = s(r[col_nr])
        fmt = s(r[col_fmt]) if col_fmt is not None and col_fmt < len(r) else ""
        gr = s(r[col_gr]) if col_gr is not None and col_gr < len(r) else ""
        farbe = s(r[col_farbe]) if col_farbe is not None and col_farbe < len(r) else ""
        for j, aufl in aufl_cols:
            p = num(r[j]) if j < len(r) else None
            if p is None:
                continue
            cello = "mit Cello" in s(cello_row[j]) if j < len(cello_row) else False
            out.append(
                {
                    "kategorie": cfg["kategorie"],
                    "produktgruppe": cfg["gruppe"],
                    "format": fmt or None,
                    "blatt": as_int(farbe),
                    "sorte": (gr or None),
                    "farbigkeit": None,
                    "spalten_key": f"{nr}_{gr}_{farbe}{'_C' if cello else ''}",
                    "auflage": aufl,
                    "preis_netto": p,
                }
            )
    return out


def main():
    path = sys.argv[1]
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    seen = {}
    for sheet in wb.sheetnames:
        if sheet in MATRIX:
            r = parse_matrix(wb[sheet], MATRIX[sheet])
        elif sheet in LONGFORM:
            r = parse_longform(wb[sheet], LONGFORM[sheet])
        else:
            continue
        seen[sheet] = len(r)
        rows.extend(r)
    json.dump({"rows": rows, "sheets": seen}, sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    main()
